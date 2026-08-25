from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import load_workbook

from amanleek_validator.application.csv_conversion import CsvConversionService
from amanleek_validator.infrastructure.excel import OpenpyxlWorkbookAdapter


class CsvConversionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.service = CsvConversionService(OpenpyxlWorkbookAdapter())

    def test_csv_is_converted_to_sheet1_and_preserves_leading_zeros(self) -> None:
        result = self.service.convert(b"Member_ID,Amount\n001,10\n")

        workbook = load_workbook(BytesIO(result.content), data_only=False)
        try:
            self.assertEqual(["sheet1"], workbook.sheetnames)
            self.assertEqual("001", workbook["sheet1"]["A2"].value)
            self.assertEqual(1, result.rows)
            self.assertEqual(2, result.columns)
        finally:
            workbook.close()

    def test_formula_like_csv_values_are_written_as_text(self) -> None:
        result = self.service.convert(b"Value\n=1+1\n")

        workbook = load_workbook(BytesIO(result.content), data_only=False)
        try:
            self.assertEqual("'=1+1", workbook["sheet1"]["A2"].value)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
