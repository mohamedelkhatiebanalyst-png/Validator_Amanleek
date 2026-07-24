from __future__ import annotations

from io import BytesIO
from zipfile import BadZipFile, ZipFile

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from amanleek_validator.domain.models import (
    SheetStructure,
    WorkbookStructure,
)
from amanleek_validator.domain.validation import normalize_header


MAX_WORKBOOK_BYTES = 100 * 1024 * 1024
MAX_UNCOMPRESSED_WORKBOOK_BYTES = 500 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 10_000
FORMULA_PREFIXES = ("=", "+", "-", "@")


class OpenpyxlWorkbookAdapter:
    def inspect(self, content: bytes) -> WorkbookStructure:
        self._validate_workbook_archive(content)
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            return tuple(self._sheet_structure(sheet) for sheet in workbook.worksheets)
        finally:
            workbook.close()

    def read_sheet(self, content: bytes, sheet_name: str) -> pd.DataFrame:
        self._validate_workbook_archive(content)
        data = pd.read_excel(
            BytesIO(content),
            sheet_name=sheet_name,
            dtype=object,
            engine="openpyxl",
        )
        data.columns = [normalize_header(column) for column in data.columns]
        return data

    def write_excel(self, data: pd.DataFrame, sheet_name: str) -> bytes:
        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            data.to_excel(writer, sheet_name=sheet_name, index=False)
        return output.getvalue()

    def write_csv(self, data: pd.DataFrame) -> bytes:
        safe_data = data.map(self._escape_csv_formula)
        return safe_data.to_csv(index=False, lineterminator="\n").encode("utf-8-sig")

    def rename_single_sheet(
        self,
        content: bytes,
        current_name: str,
        new_name: str,
    ) -> bytes:
        self._validate_workbook_archive(content)
        workbook = load_workbook(BytesIO(content))
        try:
            if workbook.sheetnames != [current_name]:
                raise ValueError("The workbook no longer has the expected single sheet")
            workbook[current_name].title = new_name
            output = BytesIO()
            workbook.save(output)
            return output.getvalue()
        finally:
            workbook.close()

    @staticmethod
    def _sheet_structure(sheet: Worksheet) -> SheetStructure:
        first_row = next(
            sheet.iter_rows(min_row=1, max_row=1, values_only=True),
            (),
        )
        return SheetStructure(
            sheet.title,
            tuple(normalize_header(value) for value in first_row),
        )

    @staticmethod
    def _escape_csv_formula(value: object) -> object:
        if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
            return "'" + value
        return value

    @staticmethod
    def _validate_workbook_archive(content: bytes) -> None:
        if not content:
            raise ValueError("The workbook is empty")
        if len(content) > MAX_WORKBOOK_BYTES:
            raise ValueError(
                f"The workbook exceeds the {MAX_WORKBOOK_BYTES // (1024 * 1024)} MB "
                "compressed-size limit"
            )
        try:
            with ZipFile(BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > MAX_ARCHIVE_MEMBERS:
                    raise ValueError("The workbook archive contains too many files")
                uncompressed_size = sum(member.file_size for member in members)
                if uncompressed_size > MAX_UNCOMPRESSED_WORKBOOK_BYTES:
                    raise ValueError(
                        "The workbook exceeds the safe uncompressed-size limit"
                    )
                if any(member.flag_bits & 0x1 for member in members):
                    raise ValueError("Encrypted workbooks are not supported")
        except BadZipFile as exc:
            raise ValueError("The uploaded file is not a valid .xlsx workbook") from exc
